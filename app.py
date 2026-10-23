import os
import logging
import hashlib
import secrets
from io import BytesIO
from functools import wraps
from datetime import datetime, date, timedelta
from calendar import monthrange
from zoneinfo import ZoneInfo

from flask import Flask, flash, redirect, render_template, request, session, url_for, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_wtf.csrf import CSRFProtect
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from flask_migrate import Migrate
from authlib.integrations.flask_client import OAuth
from sqlalchemy import func, text
from sqlalchemy.engine import URL
from werkzeug.security import check_password_hash, generate_password_hash
import requests

app = Flask(__name__)

required_settings = ('DB_USER', 'DB_PASSWORD', 'DB_HOST', 'DB_NAME', 'SECRET_KEY')
missing_settings = [name for name in required_settings if not os.environ.get(name)]
if missing_settings:
    raise RuntimeError('Faltan variables de entorno obligatorias: ' + ', '.join(missing_settings))

app.config['SQLALCHEMY_DATABASE_URI'] = URL.create(
    'mysql+pymysql', username=os.environ['DB_USER'], password=os.environ['DB_PASSWORD'],
    host=os.environ['DB_HOST'], database=os.environ['DB_NAME']
)
app.config['SECRET_KEY'] = os.environ['SECRET_KEY']
app.config['PERMANENT_SESSION_LIFETIME'] = timedelta(days=30)
app.config['SESSION_COOKIE_SECURE'] = True
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['RATELIMIT_STORAGE_URI'] = os.environ.get('RATELIMIT_STORAGE_URI', 'memory://')
app.config['GOOGLE_CLIENT_ID'] = os.environ.get('GOOGLE_CLIENT_ID')
app.config['GOOGLE_CLIENT_SECRET'] = os.environ.get('GOOGLE_CLIENT_SECRET')
# En producción usa la URL pública exacta registrada en Google Cloud.
app.config['GOOGLE_REDIRECT_URI'] = os.environ.get('GOOGLE_REDIRECT_URI')
app.config['BREVO_API_KEY'] = os.environ.get('BREVO_API_KEY')
app.config['MAIL_FROM_EMAIL'] = os.environ.get('MAIL_FROM_EMAIL')
app.config['MAIL_FROM_NAME'] = os.environ.get('MAIL_FROM_NAME', 'Láminas y Tableros Montoya')
app.config['APP_BASE_URL'] = os.environ.get('APP_BASE_URL', '').rstrip('/')
logging.basicConfig(level=os.environ.get('LOG_LEVEL', 'INFO'), format='%(asctime)s %(levelname)s %(name)s: %(message)s')
db = SQLAlchemy(app)
csrf = CSRFProtect(app)
migrate = Migrate(app, db)
limiter = Limiter(key_func=get_remote_address, app=app, default_limits=[])
oauth = OAuth(app)

if app.config['GOOGLE_CLIENT_ID'] and app.config['GOOGLE_CLIENT_SECRET']:
    oauth.register(
        name='google',
        client_id=app.config['GOOGLE_CLIENT_ID'],
        client_secret=app.config['GOOGLE_CLIENT_SECRET'],
        server_metadata_url='https://accounts.google.com/.well-known/openid-configuration',
        client_kwargs={'scope': 'openid email profile'},
    )

# ---------- MODELOS ----------

class Cliente(db.Model):
    __tablename__ = 'clientes'
    id_cliente = db.Column(db.Integer, primary_key=True)
    tipo_persona = db.Column(db.Enum('Natural', 'Juridica'), nullable=False)
    numero_identificacion = db.Column(db.String(20), unique=True, nullable=False)
    nombre_razon_social = db.Column(db.String(150), nullable=False)
    telefono = db.Column(db.String(20))
    correo = db.Column(db.String(100))
    id_usuario = db.Column(db.Integer, db.ForeignKey('usuarios.id_usuario'), unique=True, nullable=True)
    usuario = db.relationship('Usuario', backref=db.backref('cliente', uselist=False), foreign_keys=[id_usuario])


class Producto(db.Model):
    __tablename__ = 'productos'
    id_producto = db.Column(db.Integer, primary_key=True)
    tipo_estilo = db.Column(db.String(50), nullable=False)
    calibre = db.Column(db.String(10), nullable=False)
    desarrollo = db.Column(db.String(10))
    precio_por_metro = db.Column(db.Numeric(10, 2), nullable=False)


class Pedido(db.Model):
    __tablename__ = 'pedidos'
    id_pedido = db.Column(db.Integer, primary_key=True)
    id_cliente = db.Column(db.Integer, db.ForeignKey('clientes.id_cliente'), nullable=False)
    id_vendedor = db.Column(db.Integer, db.ForeignKey('usuarios.id_usuario'), nullable=True)
    fecha_pedido = db.Column(db.Date, nullable=False)
    fecha_registro = db.Column(db.DateTime, nullable=True, default=lambda: datetime.now(ZoneInfo('America/Bogota')).replace(tzinfo=None))
    estado = db.Column(db.Enum('Pendiente', 'Listo para entrega', 'Entregado'), nullable=False, default='Pendiente')
    total_pedido = db.Column(db.Numeric(10, 2), default=0)

    cliente = db.relationship('Cliente', backref='pedidos')
    vendedor = db.relationship('Usuario', foreign_keys=[id_vendedor], backref='pedidos_facturados')


class DetallePedido(db.Model):
    __tablename__ = 'detallepedido'
    id_detalle = db.Column(db.Integer, primary_key=True)
    id_pedido = db.Column(db.Integer, db.ForeignKey('pedidos.id_pedido'), nullable=False)
    id_producto = db.Column(db.Integer, db.ForeignKey('productos.id_producto'), nullable=False)
    color = db.Column(db.String(50), nullable=True)
    medida_metros = db.Column(db.Numeric(5, 2), nullable=False)
    cantidad = db.Column(db.Integer, nullable=False)
    subtotal = db.Column(db.Numeric(10, 2), nullable=False)

    pedido = db.relationship('Pedido', backref='detalles')
    producto = db.relationship('Producto')


class Usuario(db.Model):
    __tablename__ = 'usuarios'
    id_usuario = db.Column(db.Integer, primary_key=True)
    nombre = db.Column(db.String(120), nullable=False)
    correo = db.Column(db.String(120), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    rol = db.Column(db.Enum('Cliente', 'Administrador', 'Operario', 'Vendedor'), nullable=False, default='Cliente')


class SolicitudRecuperacion(db.Model):
    __tablename__ = 'solicitudes_recuperacion'
    id_solicitud = db.Column(db.Integer, primary_key=True)
    correo = db.Column(db.String(120), nullable=False)
    fecha_solicitud = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)
    estado = db.Column(db.Enum('Pendiente', 'Atendida'), nullable=False, default='Pendiente')


class TokenRecuperacionContrasena(db.Model):
    """Token de un solo uso para restablecer una contraseña.

    Solo se almacena el hash SHA-256: incluso con acceso a la base de datos no
    es posible reutilizar el enlace enviado al correo del usuario.
    """
    __tablename__ = 'tokens_recuperacion_contrasena'
    id_token = db.Column(db.Integer, primary_key=True)
    id_usuario = db.Column(db.Integer, db.ForeignKey('usuarios.id_usuario'), nullable=False, index=True)
    token_hash = db.Column(db.String(64), nullable=False, unique=True)
    fecha_creacion = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)
    fecha_expiracion = db.Column(db.DateTime, nullable=False)
    usado_en = db.Column(db.DateTime, nullable=True)
    usuario = db.relationship('Usuario', backref='tokens_recuperacion')


class Inventario(db.Model):
    __tablename__ = 'inventario'
    id_inventario = db.Column(db.Integer, primary_key=True)
    id_producto = db.Column(db.Integer, db.ForeignKey('productos.id_producto'), nullable=False)
    color = db.Column(db.String(50), nullable=False, default='General')
    cantidad_actual = db.Column(db.Integer, nullable=False, default=0)
    minimo = db.Column(db.Integer, nullable=False, default=0)
    producto = db.relationship('Producto', backref='existencias')
    __table_args__ = (db.UniqueConstraint('id_producto', 'color', name='uq_inventario_producto_color'),)


class MovimientoInventario(db.Model):
    __tablename__ = 'movimientos_inventario'
    id_movimiento = db.Column(db.Integer, primary_key=True)
    id_inventario = db.Column(db.Integer, db.ForeignKey('inventario.id_inventario'), nullable=False)
    id_usuario = db.Column(db.Integer, db.ForeignKey('usuarios.id_usuario'), nullable=True)
    tipo = db.Column(db.Enum('Entrada', 'Salida'), nullable=False)
    cantidad = db.Column(db.Integer, nullable=False)
    motivo = db.Column(db.String(180), nullable=False)
    fecha = db.Column(db.DateTime, nullable=False, default=lambda: datetime.now(ZoneInfo('America/Bogota')).replace(tzinfo=None))
    inventario = db.relationship('Inventario', backref='movimientos')


class Pago(db.Model):
    __tablename__ = 'pagos'
    id_pago = db.Column(db.Integer, primary_key=True)
    id_pedido = db.Column(db.Integer, db.ForeignKey('pedidos.id_pedido'), nullable=False)
    valor = db.Column(db.Numeric(10, 2), nullable=False)
    metodo = db.Column(db.String(50), nullable=False)
    observacion = db.Column(db.String(180))
    fecha = db.Column(db.DateTime, nullable=False, default=lambda: datetime.now(ZoneInfo('America/Bogota')).replace(tzinfo=None))
    pedido = db.relationship('Pedido', backref='pagos')


class HistorialPedido(db.Model):
    __tablename__ = 'historial_pedidos'
    id_historial = db.Column(db.Integer, primary_key=True)
    id_pedido = db.Column(db.Integer, db.ForeignKey('pedidos.id_pedido'), nullable=False)
    id_usuario = db.Column(db.Integer, db.ForeignKey('usuarios.id_usuario'), nullable=True)
    accion = db.Column(db.String(160), nullable=False)
    fecha = db.Column(db.DateTime, nullable=False, default=lambda: datetime.now(ZoneInfo('America/Bogota')).replace(tzinfo=None))
    pedido = db.relationship('Pedido', backref='historial')
    usuario = db.relationship('Usuario')


# ---------- SEGURIDAD Y ROLES ----------

def login_requerido(vista):
    @wraps(vista)
    def vista_protegida(*args, **kwargs):
        if 'usuario_id' not in session:
            flash('Inicia sesión para acceder a esta sección.', 'warning')
            return redirect(url_for('iniciar_sesion'))
        return vista(*args, **kwargs)
    return vista_protegida


def roles_requeridos(*roles):
    def decorador(vista):
        @wraps(vista)
        def vista_protegida(*args, **kwargs):
            if 'usuario_id' not in session:
                flash('Inicia sesión para acceder a esta sección.', 'warning')
                return redirect(url_for('iniciar_sesion'))
            if session.get('usuario_rol') not in roles:
                flash('No tienes permiso para acceder a esta sección.', 'danger')
                return redirect(url_for('panel'))
            return vista(*args, **kwargs)
        return vista_protegida
    return decorador


def registrar_historial(pedido, accion):
    db.session.add(HistorialPedido(id_pedido=pedido.id_pedido, id_usuario=session.get('usuario_id'), accion=accion))


def saldo_pedido(pedido):
    return max(0, float(pedido.total_pedido or 0) - sum(float(pago.valor) for pago in pedido.pagos))


def consulta_por_rol():
    consulta = Pedido.query
    if session.get('usuario_rol') == 'Vendedor':
        consulta = consulta.filter(Pedido.id_vendedor == session['usuario_id'])
    return consulta


def iniciar_sesion_usuario(usuario):
    """Crea una sesión uniforme para contraseña y proveedores externos."""
    session.clear()
    session.permanent = True
    session['usuario_id'] = usuario.id_usuario
    session['usuario_nombre'] = usuario.nombre
    session['usuario_rol'] = usuario.rol


def hash_token_recuperacion(token):
    return hashlib.sha256(token.encode('utf-8')).hexdigest()


def url_restablecimiento_publica(token):
    ruta = url_for('restablecer_contrasena', token=token)
    if app.config['APP_BASE_URL']:
        return f"{app.config['APP_BASE_URL']}{ruta}"
    return url_for('restablecer_contrasena', token=token, _external=True)


def enviar_correo_recuperacion(usuario, enlace):
    """Envía el enlace transaccional con Brevo sin exponer datos sensibles."""
    api_key = app.config['BREVO_API_KEY']
    remitente = app.config['MAIL_FROM_EMAIL']
    if not api_key or not remitente:
        app.logger.error('Recuperación de contraseña no configurada: faltan BREVO_API_KEY o MAIL_FROM_EMAIL.')
        return False

    contenido = f"""
    <div style=\"font-family:Arial,sans-serif;max-width:560px;margin:auto;color:#102f46\">
      <h2 style=\"color:#0f88c9\">Restablece tu contraseña</h2>
      <p>Hola {escape(usuario.nombre)}, recibimos una solicitud para cambiar la contraseña de tu cuenta.</p>
      <p><a href=\"{enlace}\" style=\"display:inline-block;padding:13px 20px;border-radius:8px;background:#168ecf;color:#fff;text-decoration:none;font-weight:bold\">Crear nueva contraseña</a></p>
      <p>Este enlace vence en 15 minutos y solo puede usarse una vez.</p>
      <p>Si no solicitaste este cambio, puedes ignorar este correo.</p>
    </div>
    """
    payload = {
        'sender': {'email': remitente, 'name': app.config['MAIL_FROM_NAME']},
        'to': [{'email': usuario.correo, 'name': usuario.nombre}],
        'subject': 'Restablece tu contraseña | Láminas y Tableros Montoya',
        'htmlContent': contenido,
    }
    try:
        respuesta = requests.post(
            'https://api.brevo.com/v3/smtp/email',
            headers={'api-key': api_key, 'accept': 'application/json', 'content-type': 'application/json'},
            json=payload,
            timeout=12,
        )
        if respuesta.ok:
            app.logger.info('Correo de recuperación enviado para usuario=%s.', usuario.id_usuario)
            return True
        app.logger.warning('Brevo rechazó un correo de recuperación: estado=%s.', respuesta.status_code)
    except requests.RequestException:
        app.logger.exception('No fue posible comunicarse con Brevo para recuperación de contraseña.')
    return False


@app.context_processor
def variables_de_acceso():
    """Evita mostrar el acceso de Google si faltan sus credenciales."""
    return {'google_login_habilitado': oauth.create_client('google') is not None}


@app.errorhandler(429)
def demasiadas_peticiones(error):
    flash('Has solicitado varios enlaces recientemente. Espera unos minutos antes de pedir otro.', 'warning')
    return redirect(url_for('recuperar_contrasena'))


@app.errorhandler(500)
def error_interno(error):
    app.logger.exception('Error no controlado: %s', error)
    return render_template('error.html'), 500

# ---------- RUTAS: CLIENTES ----------

@app.route('/')
@app.route('/empresa')
def pagina_empresa():
    return render_template('empresa.html')


# ---------- ACCESO DE USUARIOS ----------

@app.route('/configurar-admin-inicial', methods=['GET', 'POST'])
def configurar_admin_inicial():
    if Usuario.query.filter_by(rol='Administrador').first():
        flash('El administrador inicial ya fue configurado. Inicia sesión o crea una cuenta de cliente.', 'info')
        return redirect(url_for('iniciar_sesion'))

    if request.method == 'POST':
        usuario = Usuario(
            nombre=request.form['nombre'].strip(),
            correo=request.form['correo'].strip().lower(),
            password_hash=generate_password_hash(request.form['password']),
            rol='Administrador'
        )
        db.session.add(usuario)
        db.session.commit()
        flash('Administrador creado. Ya puedes iniciar sesión.', 'success')
        return redirect(url_for('iniciar_sesion'))
    return render_template('acceso.html', modo='administrador')


@app.route('/crear-cuenta', methods=['GET', 'POST'])
def crear_cuenta():
    if request.method == 'POST':
        correo = request.form['correo'].strip().lower()
        if Usuario.query.filter_by(correo=correo).first():
            flash('Ya existe una cuenta con ese correo.', 'danger')
            return redirect(url_for('crear_cuenta'))
        if len(request.form['password']) < 8:
            flash('La contraseña debe tener al menos 8 caracteres.', 'danger')
            return redirect(url_for('crear_cuenta'))

        identificacion = request.form['numero_identificacion'].strip()
        if Cliente.query.filter_by(numero_identificacion=identificacion).first():
            flash('Ya existe un cliente con ese número de identificación.', 'danger')
            return redirect(url_for('crear_cuenta'))
        usuario = Usuario(
            nombre=request.form['nombre'].strip(), correo=correo,
            password_hash=generate_password_hash(request.form['password']), rol='Cliente'
        )
        db.session.add(usuario)
        db.session.flush()
        db.session.add(Cliente(
            tipo_persona=request.form['tipo_persona'], numero_identificacion=identificacion,
            nombre_razon_social=usuario.nombre, telefono=request.form['telefono'].strip(),
            correo=correo, id_usuario=usuario.id_usuario
        ))
        db.session.commit()
        flash('Tu cuenta fue creada. Ya puedes iniciar sesión.', 'success')
        return redirect(url_for('iniciar_sesion'))
    return render_template('acceso.html', modo='registro')


@app.route('/iniciar-sesion', methods=['GET', 'POST'])
@limiter.limit('5 per minute')
def iniciar_sesion():
    if request.method == 'POST':
        usuario = Usuario.query.filter_by(correo=request.form['correo'].strip().lower()).first()
        if usuario and check_password_hash(usuario.password_hash, request.form['password']):
            iniciar_sesion_usuario(usuario)
            return redirect(url_for('panel'))
        app.logger.warning('Intento de inicio de sesión fallido para %s desde %s', request.form['correo'].strip().lower(), get_remote_address())
        flash('Correo o contraseña incorrectos.', 'danger')
    return render_template('acceso.html', modo='login')


@app.route('/iniciar-sesion/google')
@limiter.limit('10 per minute')
def iniciar_sesion_google():
    """Envía a Google solo para autenticar una cuenta ya registrada."""
    google = oauth.create_client('google')
    if google is None:
        flash('El acceso con Google aún no está configurado.', 'warning')
        return redirect(url_for('iniciar_sesion'))

    redirect_uri = app.config['GOOGLE_REDIRECT_URI'] or url_for('autorizar_google', _external=True)
    return google.authorize_redirect(redirect_uri, prompt='select_account')


@app.route('/autorizar/google')
@limiter.limit('10 per minute')
def autorizar_google():
    """Valida la identidad devuelta por Google e inicia la sesión local."""
    google = oauth.create_client('google')
    if google is None:
        flash('El acceso con Google aún no está configurado.', 'warning')
        return redirect(url_for('iniciar_sesion'))

    if request.args.get('error'):
        flash('No se completó el acceso con Google. Puedes intentarlo nuevamente.', 'warning')
        return redirect(url_for('iniciar_sesion'))

    try:
        token = google.authorize_access_token()
        datos_google = token.get('userinfo') or {}
    except Exception:
        app.logger.exception('No fue posible validar el inicio de sesión con Google.')
        flash('No fue posible validar tu cuenta de Google. Inténtalo nuevamente.', 'danger')
        return redirect(url_for('iniciar_sesion'))

    correo = str(datos_google.get('email') or '').strip().lower()
    correo_verificado = datos_google.get('email_verified') is True or str(datos_google.get('email_verified')).lower() == 'true'
    if not correo or not correo_verificado:
        flash('Google no confirmó un correo electrónico verificado para esta cuenta.', 'danger')
        return redirect(url_for('iniciar_sesion'))

    usuario = Usuario.query.filter_by(correo=correo).first()
    if usuario is None:
        # No se crean clientes incompletos: identificación y teléfono siguen siendo obligatorios.
        flash('Tu correo de Google aún no tiene una cuenta. Regístrate una vez y luego podrás entrar con Google.', 'info')
        return redirect(url_for('crear_cuenta'))

    iniciar_sesion_usuario(usuario)
    return redirect(url_for('panel'))


@app.route('/recuperar-contrasena', methods=['GET', 'POST'])
@limiter.limit('3 per 15 minutes', methods=['POST'])
def recuperar_contrasena():
    if request.method == 'POST':
        correo = request.form['correo'].strip().lower()
        usuario = Usuario.query.filter_by(correo=correo).first()
        if usuario:
            ahora = datetime.utcnow()
            # Un nuevo enlace anula cualquiera anterior que aún no se hubiera usado.
            TokenRecuperacionContrasena.query.filter(
                TokenRecuperacionContrasena.id_usuario == usuario.id_usuario,
                TokenRecuperacionContrasena.usado_en.is_(None),
            ).update({'usado_en': ahora}, synchronize_session=False)
            token = secrets.token_urlsafe(32)
            registro = TokenRecuperacionContrasena(
                id_usuario=usuario.id_usuario,
                token_hash=hash_token_recuperacion(token),
                fecha_expiracion=ahora + timedelta(minutes=15),
            )
            db.session.add(registro)
            db.session.commit()
            if not enviar_correo_recuperacion(usuario, url_restablecimiento_publica(token)):
                # No conservamos un token que el usuario no pudo recibir.
                db.session.delete(registro)
                db.session.commit()
        # Mensaje genérico: no revela si el correo existe en la aplicación.
        flash('Si el correo está registrado, recibirás un enlace de recuperación. Revisa también tu carpeta de spam.', 'success')
        return redirect(url_for('iniciar_sesion'))
    return render_template('acceso.html', modo='recuperar')


@app.route('/restablecer-contrasena/<token>', methods=['GET', 'POST'])
@limiter.limit('10 per hour')
def restablecer_contrasena(token):
    registro = TokenRecuperacionContrasena.query.filter_by(
        token_hash=hash_token_recuperacion(token)
    ).first()
    ahora = datetime.utcnow()
    if not registro or registro.usado_en or registro.fecha_expiracion < ahora:
        flash('Este enlace ya fue usado o venció. Solicita uno nuevo.', 'warning')
        return redirect(url_for('recuperar_contrasena'))

    if request.method == 'POST':
        password = request.form.get('password', '')
        confirmar_password = request.form.get('confirmar_password', '')
        if len(password) < 8:
            flash('La contraseña debe tener al menos 8 caracteres.', 'danger')
            return render_template('acceso.html', modo='restablecer', token=token)
        if password != confirmar_password:
            flash('Las contraseñas no coinciden.', 'danger')
            return render_template('acceso.html', modo='restablecer', token=token)

        registro.usuario.password_hash = generate_password_hash(password)
        registro.usado_en = ahora
        TokenRecuperacionContrasena.query.filter(
            TokenRecuperacionContrasena.id_usuario == registro.id_usuario,
            TokenRecuperacionContrasena.usado_en.is_(None),
            TokenRecuperacionContrasena.id_token != registro.id_token,
        ).update({'usado_en': ahora}, synchronize_session=False)
        db.session.commit()
        flash('Tu contraseña fue actualizada. Ya puedes iniciar sesión.', 'success')
        return redirect(url_for('iniciar_sesion'))

    return render_template('acceso.html', modo='restablecer', token=token)


@app.route('/cerrar-sesion')
def cerrar_sesion():
    session.clear()
    flash('Sesión cerrada correctamente.', 'info')
    return redirect(url_for('pagina_empresa'))


@app.route('/panel')
@login_requerido
def panel():
    rol = session['usuario_rol']
    datos = {}
    resumen_mensual = None
    ejecutivo = None
    if rol == 'Administrador':
        datos = {'Clientes registrados': Cliente.query.count(), 'Productos activos': Producto.query.count(), 'Pedidos registrados': Pedido.query.count()}
        inicio_mes = date.today().replace(day=1)
        pedidos_mes = Pedido.query.filter(Pedido.fecha_pedido >= inicio_mes).all()
        pagos_mes = Pago.query.filter(Pago.fecha >= datetime.combine(inicio_mes, datetime.min.time())).all()
        pedidos_todos = Pedido.query.all()
        alertas_inventario = Inventario.query.filter(Inventario.cantidad_actual <= Inventario.minimo).count()
        proximas_entregas = Pedido.query.filter(
            Pedido.fecha_pedido.between(date.today(), date.today() + timedelta(days=7)),
            Pedido.estado != 'Entregado'
        ).order_by(Pedido.fecha_pedido.asc()).limit(5).all()
        cobro_pendiente = sum(saldo_pedido(pedido) for pedido in pedidos_todos)
        resumen_mensual = {
            'Ventas del mes': sum(float(p.total_pedido or 0) for p in pedidos_mes),
            'Entregados': sum(p.estado == 'Entregado' for p in pedidos_mes),
            'Pendientes': sum(p.estado == 'Pendiente' for p in pedidos_mes),
        }
        ejecutivo = {
            'cobrado_mes': sum(float(pago.valor or 0) for pago in pagos_mes),
            'por_cobrar': cobro_pendiente,
            'alertas_inventario': alertas_inventario,
            'proximas_entregas': proximas_entregas,
            'pagos_recientes': Pago.query.order_by(Pago.fecha.desc()).limit(5).all(),
        }
    elif rol == 'Operario':
        datos = {'Pedidos pendientes': Pedido.query.filter_by(estado='Pendiente').count(), 'Listos para entrega': Pedido.query.filter_by(estado='Listo para entrega').count(), 'Pedidos entregados': Pedido.query.filter_by(estado='Entregado').count()}
    elif rol == 'Vendedor':
        pedidos_vendedor = Pedido.query.filter_by(id_vendedor=session['usuario_id'])
        datos = {'Pedidos registrados': pedidos_vendedor.count(), 'Pendientes': pedidos_vendedor.filter_by(estado='Pendiente').count(), 'Listos para entrega': pedidos_vendedor.filter_by(estado='Listo para entrega').count()}
    else:
        cliente = Cliente.query.filter_by(id_usuario=session['usuario_id']).first()
        pedidos_cliente = Pedido.query.filter_by(id_cliente=cliente.id_cliente).count() if cliente else 0
        datos = {'Mis pedidos': pedidos_cliente, 'Estado de cuenta': 'Activo', 'Atención': 'Contáctanos para cotizar'}
    return render_template('panel.html', datos=datos, rol=rol, resumen_mensual=resumen_mensual, ejecutivo=ejecutivo)


@app.route('/reportes')
@roles_requeridos('Administrador', 'Operario', 'Vendedor')
def reportes():
    """Indicadores comerciales y operativos, respetando el alcance de cada perfil."""
    rol = session['usuario_rol']
    periodo = request.args.get('periodo', '30')
    periodos = {'7': 'Últimos 7 días', '30': 'Últimos 30 días', '90': 'Últimos 90 días', 'todos': 'Todo el historial'}
    if periodo not in periodos:
        periodo = '30'

    fecha_inicio = request.args.get('fecha_inicio', '')
    fecha_fin = request.args.get('fecha_fin', '')
    consulta = consulta_por_rol()
    etiqueta_periodo = periodos[periodo]
    if rol in ('Administrador', 'Vendedor') and fecha_inicio and fecha_fin:
        try:
            inicio = date.fromisoformat(fecha_inicio)
            fin = date.fromisoformat(fecha_fin)
            if inicio <= fin:
                consulta = consulta.filter(Pedido.fecha_pedido.between(inicio, fin))
                periodo = 'rango'
                etiqueta_periodo = f'{inicio.strftime("%d/%m/%Y")} — {fin.strftime("%d/%m/%Y")}'
            else:
                fecha_inicio = fecha_fin = ''
        except ValueError:
            fecha_inicio = fecha_fin = ''
    elif periodo != 'todos':
        consulta = consulta.filter(Pedido.fecha_pedido >= date.today() - timedelta(days=int(periodo) - 1))
    pedidos_periodo = consulta.order_by(Pedido.fecha_pedido.asc()).all()

    total_ventas = sum(float(p.total_pedido or 0) for p in pedidos_periodo)
    total_pedidos = len(pedidos_periodo)
    por_estado = {estado: 0 for estado in ('Pendiente', 'Listo para entrega', 'Entregado')}
    ventas_por_dia = {}
    pedidos_por_dia = {}
    productos = {}
    clientes = {}
    for pedido in pedidos_periodo:
        por_estado[pedido.estado] = por_estado.get(pedido.estado, 0) + 1
        dia = pedido.fecha_pedido.strftime('%d %b')
        ventas_por_dia[dia] = ventas_por_dia.get(dia, 0) + float(pedido.total_pedido or 0)
        pedidos_por_dia[dia] = pedidos_por_dia.get(dia, 0) + 1
        nombre_cliente = pedido.cliente.nombre_razon_social
        clientes[nombre_cliente] = clientes.get(nombre_cliente, 0) + float(pedido.total_pedido or 0)
        for detalle in pedido.detalles:
            nombre_producto = f'{detalle.producto.tipo_estilo} · Cal. {detalle.producto.calibre}'
            actual = productos.setdefault(nombre_producto, {'cantidad': 0, 'metros': 0})
            actual['cantidad'] += detalle.cantidad
            actual['metros'] += float(detalle.medida_metros) * detalle.cantidad

    top_productos = sorted(productos.items(), key=lambda item: item[1]['metros'], reverse=True)[:5]
    top_clientes = sorted(clientes.items(), key=lambda item: item[1], reverse=True)[:5]
    metricas = {'Pedidos': total_pedidos, 'Metros solicitados': sum(item['metros'] for item in productos.values())}
    if rol != 'Operario':
        metricas.update({'Valor total': total_ventas, 'Ticket promedio': total_ventas / total_pedidos if total_pedidos else 0})
    return render_template(
        'reportes.html', rol=rol, periodo=periodo, periodos=periodos, etiqueta_periodo=etiqueta_periodo,
        fecha_inicio=fecha_inicio, fecha_fin=fecha_fin, metricas=metricas,
        por_estado=por_estado, ventas_labels=list(ventas_por_dia.keys()),
        ventas_data=list(pedidos_por_dia.values()) if rol == 'Operario' else list(ventas_por_dia.values()),
        top_productos=top_productos, top_clientes=top_clientes
    )


@app.route('/inventario', methods=['GET', 'POST'])
@roles_requeridos('Administrador', 'Operario')
def inventario():
    if request.method == 'POST':
        if session['usuario_rol'] != 'Administrador':
            flash('Solo Administración puede registrar movimientos de inventario.', 'danger')
            return redirect(url_for('inventario'))
        item = Inventario.query.get_or_404(request.form.get('id_inventario', type=int))
        cantidad = request.form.get('cantidad', type=int)
        if not cantidad or cantidad <= 0:
            flash('Indica una cantidad válida.', 'danger')
            return redirect(url_for('inventario'))
        cliente_venta = request.form.get('cliente_venta', '').strip()
        observacion = request.form.get('observacion', '').strip()
        if not cliente_venta:
            flash('Indica el cliente o destino de la venta.', 'danger')
            return redirect(url_for('inventario'))
        if item.cantidad_actual < cantidad:
            flash('No hay existencias suficientes para registrar esta venta.', 'danger')
            return redirect(url_for('inventario'))
        item.cantidad_actual -= cantidad
        motivo = f'Venta directa · {cliente_venta}'
        if observacion:
            motivo += f' · {observacion}'
        db.session.add(MovimientoInventario(id_inventario=item.id_inventario, id_usuario=session['usuario_id'], tipo='Salida', cantidad=cantidad, motivo=motivo[:180]))
        db.session.commit()
        flash('Venta directa registrada e inventario actualizado.', 'success')
        return redirect(url_for('inventario'))
    productos = Producto.query.order_by(Producto.tipo_estilo).all()
    items = Inventario.query.join(Producto).order_by(Producto.tipo_estilo, Inventario.color).all()
    alertas_inventario = [item for item in items if item.cantidad_actual <= item.minimo]
    ventas_directas = MovimientoInventario.query.filter(
        MovimientoInventario.tipo == 'Salida',
        MovimientoInventario.motivo.like('Venta directa ·%')
    ).order_by(MovimientoInventario.fecha.desc()).limit(8).all()
    return render_template('inventario.html', items=items, productos=productos, alertas_inventario=alertas_inventario,
                           ventas_directas=ventas_directas)


@app.route('/inventario/nuevo', methods=['POST'])
@roles_requeridos('Administrador')
def nuevo_item_inventario():
    producto_id = request.form.get('id_producto', type=int)
    color = request.form['color'].strip() or 'General'
    if Inventario.query.filter_by(id_producto=producto_id, color=color).first():
        flash('Ya existe inventario para ese producto y color.', 'warning')
    else:
        db.session.add(Inventario(id_producto=producto_id, color=color, cantidad_actual=request.form.get('cantidad_actual', type=int) or 0, minimo=request.form.get('minimo', type=int) or 0))
        db.session.commit()
        flash('Referencia agregada al inventario.', 'success')
    return redirect(url_for('inventario'))


@app.route('/calendario')
@roles_requeridos('Administrador', 'Operario', 'Vendedor')
def calendario():
    hoy = date.today()
    anio = request.args.get('anio', hoy.year, type=int)
    mes = request.args.get('mes', hoy.month, type=int)
    if mes < 1 or mes > 12:
        mes = hoy.month
    primer_dia = date(anio, mes, 1)
    ultimo_dia = date(anio, mes, monthrange(anio, mes)[1])
    pedidos = consulta_por_rol().filter(Pedido.fecha_pedido.between(primer_dia, ultimo_dia)).order_by(Pedido.fecha_pedido).all()
    pedidos_por_fecha = {}
    for pedido in pedidos:
        pedidos_por_fecha.setdefault(pedido.fecha_pedido.isoformat(), []).append(pedido)
    anterior = (primer_dia - timedelta(days=1)).replace(day=1)
    siguiente = (ultimo_dia + timedelta(days=1)).replace(day=1)
    return render_template('calendario.html', pedidos_por_fecha=pedidos_por_fecha, hoy=hoy, mes=mes, anio=anio, primer_dia=primer_dia, dias_mes=ultimo_dia.day, inicio_semana=primer_dia.weekday(), anterior=anterior, siguiente=siguiente)


@app.route('/pagos', methods=['GET', 'POST'])
@roles_requeridos('Administrador')
def pagos():
    if request.method == 'POST':
        pedido = Pedido.query.get_or_404(request.form.get('id_pedido', type=int))
        valor = request.form.get('valor', type=float)
        if not valor or valor <= 0 or valor > saldo_pedido(pedido):
            flash('El pago debe ser mayor a cero y no superar el saldo pendiente.', 'danger')
        else:
            db.session.add(Pago(id_pedido=pedido.id_pedido, valor=valor, metodo=request.form['metodo'], observacion=request.form.get('observacion', '').strip()))
            registrar_historial(pedido, f'Pago registrado por ${valor:,.0f}')
            db.session.commit()
            flash('Pago registrado correctamente.', 'success')
        return redirect(url_for('pagos'))
    pedidos = Pedido.query.order_by(Pedido.fecha_pedido.desc()).all()
    saldos_pendientes = [pedido for pedido in pedidos if saldo_pedido(pedido) > 0]
    total_por_cobrar = sum(saldo_pedido(pedido) for pedido in saldos_pendientes)
    pagos_recientes = Pago.query.order_by(Pago.fecha.desc()).limit(12).all()
    return render_template('pagos.html', pedidos=pedidos, saldo_pedido=saldo_pedido,
                           saldos_pendientes=saldos_pendientes, total_por_cobrar=total_por_cobrar,
                           pagos_recientes=pagos_recientes)


@app.route('/pagos/<int:id>/comprobante.pdf')
@roles_requeridos('Administrador')
def comprobante_pago(id):
    """Genera un recibo interno de pago; no equivale a factura electrónica."""
    pago = Pago.query.get_or_404(id)
    pedido = pago.pedido
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    salida = BytesIO()
    documento = SimpleDocTemplate(salida, pagesize=letter, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42)
    estilos = getSampleStyleSheet()
    estilos.add(ParagraphStyle(name='ReceiptNote', parent=estilos['Normal'], textColor=colors.HexColor('#63798C'), fontSize=8, leading=11))
    encabezado = Table([[Paragraph('<b>LÁMINAS Y TABLEROS MONTOYA</b><br/><font size="9">COMPROBANTE DE PAGO</font>', estilos['Title']), Paragraph(f'<b>RECIBO #{pago.id_pago:05d}</b><br/><font size="9">{pago.fecha.strftime("%d/%m/%Y %H:%M")}</font>', estilos['Normal'])]], colWidths=[4.65 * inch, 2.15 * inch])
    encabezado.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#092B52')), ('TEXTCOLOR', (0, 0), (-1, -1), colors.white), ('PADDING', (0, 0), (-1, -1), 15), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('ALIGN', (1, 0), (1, 0), 'RIGHT')]))
    saldo_actual = saldo_pedido(pedido)
    detalle = [
        ['Cliente', pedido.cliente.nombre_razon_social],
        ['Pedido relacionado', f'#{pedido.id_pedido:05d} · Entrega {pedido.fecha_pedido.strftime("%d/%m/%Y")}'],
        ['Método de pago', pago.metodo],
        ['Observación', pago.observacion or 'Sin observaciones'],
        ['Valor recibido', f'${float(pago.valor):,.0f} COP'],
        ['Saldo actual del pedido', f'${saldo_actual:,.0f} COP'],
    ]
    tabla = Table(detalle, colWidths=[2.05 * inch, 4.75 * inch])
    tabla.setStyle(TableStyle([('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#E9F6FC')), ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#0D4D86')), ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'), ('GRID', (0, 0), (-1, -1), .35, colors.HexColor('#CFE4F0')), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('PADDING', (0, 0), (-1, -1), 10)]))
    documento.build([encabezado, Spacer(1, 18), Paragraph('Detalle del recibo', estilos['Heading2']), Spacer(1, 8), tabla, Spacer(1, 20), Paragraph('Este documento es un comprobante interno de pago y no reemplaza una factura electrónica.', estilos['ReceiptNote'])])
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=f'comprobante_pago_{pago.id_pago:05d}.pdf', mimetype='application/pdf')


@app.route('/clientes/<int:id>/estado-cuenta.pdf')
@roles_requeridos('Administrador')
def estado_cuenta_cliente(id):
    """Genera el estado de cuenta interno de un cliente, usando pedidos y pagos existentes."""
    cliente = Cliente.query.get_or_404(id)
    pedidos_cliente = Pedido.query.filter_by(id_cliente=cliente.id_cliente).order_by(Pedido.fecha_pedido.desc(), Pedido.id_pedido.desc()).all()
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    total_facturado = sum(float(pedido.total_pedido or 0) for pedido in pedidos_cliente)
    total_pagado = sum(sum(float(pago.valor or 0) for pago in pedido.pagos) for pedido in pedidos_cliente)
    total_saldo = sum(saldo_pedido(pedido) for pedido in pedidos_cliente)
    salida = BytesIO()
    documento = SimpleDocTemplate(salida, pagesize=letter, rightMargin=34, leftMargin=34, topMargin=42, bottomMargin=42)
    estilos = getSampleStyleSheet()
    estilos.add(ParagraphStyle(name='AccountNote', parent=estilos['Normal'], textColor=colors.HexColor('#63798C'), fontSize=8, leading=11))
    encabezado = Table([[Paragraph('<b>LÁMINAS Y TABLEROS MONTOYA</b><br/><font size="9">ESTADO DE CUENTA</font>', estilos['Title']), Paragraph(f'<b>GENERADO</b><br/><font size="9">{datetime.now(ZoneInfo("America/Bogota")).strftime("%d/%m/%Y %H:%M")}</font>', estilos['Normal'])]], colWidths=[4.7 * inch, 2.1 * inch])
    encabezado.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#092B52')), ('TEXTCOLOR', (0, 0), (-1, -1), colors.white), ('PADDING', (0, 0), (-1, -1), 15), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('ALIGN', (1, 0), (1, 0), 'RIGHT')]))
    resumen = Table([['CLIENTE', cliente.nombre_razon_social], ['IDENTIFICACIÓN', cliente.numero_identificacion], ['TOTAL FACTURADO', f'${total_facturado:,.0f} COP'], ['TOTAL PAGADO', f'${total_pagado:,.0f} COP'], ['SALDO PENDIENTE', f'${total_saldo:,.0f} COP']], colWidths=[2.05 * inch, 4.75 * inch])
    resumen.setStyle(TableStyle([('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#E9F6FC')), ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#0D4D86')), ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'), ('GRID', (0, 0), (-1, -1), .35, colors.HexColor('#CFE4F0')), ('PADDING', (0, 0), (-1, -1), 8)]))
    filas = [['Pedido', 'Entrega', 'Estado', 'Total', 'Pagado', 'Saldo']]
    for pedido in pedidos_cliente:
        pagado = float(pedido.total_pedido or 0) - saldo_pedido(pedido)
        filas.append([f'#{pedido.id_pedido:05d}', pedido.fecha_pedido.strftime('%d/%m/%Y'), pedido.estado, f'${float(pedido.total_pedido or 0):,.0f}', f'${pagado:,.0f}', f'${saldo_pedido(pedido):,.0f}'])
    detalle = Table(filas, repeatRows=1, colWidths=[.72 * inch, .86 * inch, 1.45 * inch, 1.2 * inch, 1.2 * inch, 1.2 * inch])
    detalle.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1279C9')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#CFE4F0')), ('FONTSIZE', (0, 0), (-1, -1), 8), ('ALIGN', (3, 1), (-1, -1), 'RIGHT'), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5FBFE')]), ('PADDING', (0, 0), (-1, -1), 7)]))
    documento.build([encabezado, Spacer(1, 16), resumen, Spacer(1, 20), Paragraph('Detalle de pedidos', estilos['Heading2']), Spacer(1, 8), detalle, Spacer(1, 14), Paragraph('Estado de cuenta interno sujeto a validación administrativa.', estilos['AccountNote'])])
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=f'estado_cuenta_cliente_{cliente.id_cliente}.pdf', mimetype='application/pdf')


@app.route('/pedidos/<int:id>/historial')
@roles_requeridos('Administrador', 'Operario', 'Vendedor')
def historial_pedido(id):
    pedido = Pedido.query.get_or_404(id)
    if session['usuario_rol'] == 'Vendedor' and pedido.id_vendedor != session['usuario_id']:
        flash('Solo puedes consultar el historial de tus pedidos.', 'danger')
        return redirect(url_for('ver_pedidos'))
    return render_template('historial_pedido.html', pedido=pedido)


@app.route('/reportes/exportar/<formato>')
@roles_requeridos('Administrador', 'Operario', 'Vendedor')
def exportar_reporte(formato):
    periodo = request.args.get('periodo', '30')
    if periodo not in ('7', '30', '90', 'todos'):
        periodo = '30'
    fecha_inicio = request.args.get('fecha_inicio', '')
    fecha_fin = request.args.get('fecha_fin', '')
    consulta = consulta_por_rol()
    if session['usuario_rol'] in ('Administrador', 'Vendedor') and fecha_inicio and fecha_fin:
        try:
            inicio = date.fromisoformat(fecha_inicio)
            fin = date.fromisoformat(fecha_fin)
            if inicio <= fin:
                consulta = consulta.filter(Pedido.fecha_pedido.between(inicio, fin))
            else:
                fecha_inicio = fecha_fin = ''
        except ValueError:
            fecha_inicio = fecha_fin = ''
    if not (fecha_inicio and fecha_fin) and periodo != 'todos':
        consulta = consulta.filter(Pedido.fecha_pedido >= date.today() - timedelta(days=int(periodo) - 1))
    pedidos = consulta.order_by(Pedido.fecha_pedido.desc()).all()
    filas = [(p.id_pedido, p.cliente.nombre_razon_social, p.fecha_pedido, p.estado, float(p.total_pedido or 0)) for p in pedidos]
    mostrar_valor = session['usuario_rol'] != 'Operario'
    encabezados = ['Pedido', 'Cliente', 'Entrega', 'Estado'] + (['Valor'] if mostrar_valor else [])
    if formato == 'excel':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        libro = Workbook(); hoja = libro.active; hoja.title = 'Reporte de pedidos'
        hoja.append(encabezados)
        for celda in hoja[1]:
            celda.font = Font(bold=True, color='FFFFFF'); celda.fill = PatternFill('solid', fgColor='1279C9')
        for fila in filas:
            hoja.append(fila if mostrar_valor else fila[:-1])
        for columna in hoja.columns:
            hoja.column_dimensions[columna[0].column_letter].width = min(max(len(str(c.value or '')) for c in columna) + 3, 35)
        salida = BytesIO(); libro.save(salida); salida.seek(0)
        return send_file(salida, as_attachment=True, download_name='reporte_pedidos.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    if formato == 'pdf':
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.units import inch
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
        salida = BytesIO(); documento = SimpleDocTemplate(salida, pagesize=letter, rightMargin=38, leftMargin=38, topMargin=44, bottomMargin=38)
        estilos = getSampleStyleSheet(); total = sum(fila[4] for fila in filas)
        estados = {estado: sum(1 for fila in filas if fila[3] == estado) for estado in ('Pendiente', 'Listo para entrega', 'Entregado')}
        titulo = 'REPORTE COMERCIAL' if mostrar_valor else 'REPORTE OPERATIVO'
        cabecera = Table([[Paragraph('<b>LÁMINAS Y TABLEROS MONTOYA</b><br/><font size="9">' + titulo + '</font>', estilos['Title']), Paragraph('<b>Generado</b><br/><font size="9">' + datetime.now(ZoneInfo('America/Bogota')).strftime('%d/%m/%Y %H:%M') + '</font>', estilos['Normal'])]], colWidths=[4.7*inch, 2.2*inch])
        cabecera.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#092B52')),('TEXTCOLOR',(0,0),(-1,-1),colors.white),('PADDING',(0,0),(-1,-1),16),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ALIGN',(1,0),(1,0),'RIGHT')]))
        tarjetas = [['PEDIDOS', str(len(filas)), 'PENDIENTES', str(estados['Pendiente'])]]
        if mostrar_valor:
            tarjetas[0] += ['VALOR TOTAL', f'${total:,.0f}']
        else:
            tarjetas[0] += ['LISTOS PARA ENTREGA', str(estados['Listo para entrega'])]
        kpis = Table(tarjetas, colWidths=([1.15*inch, .8*inch] * 3))
        kpis.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#E9F6FC')),('TEXTCOLOR',(0,0),(-1,-1),colors.HexColor('#0D4D86')),('FONTNAME',(0,0),(-1,-1),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),9),('BOX',(0,0),(-1,-1),.5,colors.HexColor('#CFE4F0')),('INNERGRID',(0,0),(-1,-1),.5,colors.HexColor('#CFE4F0')),('ALIGN',(0,0),(-1,-1),'CENTER'),('PADDING',(0,0),(-1,-1),12)]))
        data = [encabezados]
        for fila in filas:
            data.append([fila[0], fila[1], fila[2].strftime('%d/%m/%Y'), fila[3]] + ([f'${fila[4]:,.0f}'] if mostrar_valor else []))
        tabla = Table(data, repeatRows=1, colWidths=[52, 185, 76, 105] + ([78] if mostrar_valor else []))
        tabla.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1279C9')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#CFE4F0')), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('FONTSIZE', (0, 0), (-1, -1), 8), ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F5FBFE')]), ('BOTTOMPADDING', (0, 0), (-1, -1), 8), ('TOPPADDING', (0, 0), (-1, -1), 8)]))
        documento.build([cabecera, Spacer(1, 16), kpis, Spacer(1, 20), Paragraph('Detalle de pedidos', estilos['Heading2']), Spacer(1, 8), tabla, Spacer(1, 16), Paragraph('Documento generado por el sistema de gestión de LÁMINAS Y TABLEROS MONTOYA.', estilos['Normal'])])
        salida.seek(0)
        return send_file(salida, as_attachment=True, download_name='reporte_pedidos.pdf', mimetype='application/pdf')
    return redirect(url_for('reportes'))


@app.route('/mis-pedidos')
@roles_requeridos('Cliente')
def mis_pedidos():
    cliente = Cliente.query.filter_by(id_usuario=session['usuario_id']).first()
    pedidos = Pedido.query.filter_by(id_cliente=cliente.id_cliente).order_by(Pedido.fecha_pedido.desc()).all() if cliente else []
    return render_template('mis_pedidos.html', pedidos=pedidos, cliente_registrado=cliente is not None)


@app.route('/usuarios')
@roles_requeridos('Administrador')
def usuarios():
    solicitudes = SolicitudRecuperacion.query.filter_by(estado='Pendiente').order_by(SolicitudRecuperacion.fecha_solicitud.desc()).all()
    return render_template('usuarios.html', usuarios=Usuario.query.order_by(Usuario.nombre).all(), solicitudes=solicitudes)


@app.route('/usuarios/<int:id>/rol', methods=['POST'])
@roles_requeridos('Administrador')
def actualizar_rol(id):
    usuario = Usuario.query.get_or_404(id)
    nuevo_rol = request.form['rol']
    if nuevo_rol not in ('Cliente', 'Administrador', 'Operario', 'Vendedor'):
        flash('Rol no válido.', 'danger')
    elif usuario.id_usuario == session['usuario_id'] and nuevo_rol != 'Administrador':
        flash('No puedes quitarte tu propio rol de administrador.', 'danger')
    else:
        app.logger.info('Cambio de rol: usuario=%s, nuevo_rol=%s, administrador=%s', usuario.id_usuario, nuevo_rol, session['usuario_id'])
        usuario.rol = nuevo_rol
        db.session.commit()
        flash(f'El rol de {usuario.nombre} fue actualizado.', 'success')
    return redirect(url_for('usuarios'))


@app.route('/usuarios/<int:id>/eliminar', methods=['POST'])
@roles_requeridos('Administrador')
def eliminar_usuario(id):
    usuario = Usuario.query.get_or_404(id)
    if usuario.id_usuario == session['usuario_id']:
        flash('No puedes eliminar tu propia cuenta mientras tienes la sesión activa.', 'danger')
    elif usuario.rol == 'Administrador' and Usuario.query.filter_by(rol='Administrador').count() <= 1:
        flash('No puedes eliminar al único administrador del sistema.', 'danger')
    elif Pedido.query.filter_by(id_vendedor=usuario.id_usuario).first():
        flash(f'No puedes eliminar a {usuario.nombre} porque tiene pedidos registrados. Conserva la trazabilidad o reasigna sus pedidos.', 'danger')
    else:
        app.logger.info('Eliminación de usuario: usuario=%s, administrador=%s', usuario.id_usuario, session['usuario_id'])
        db.session.delete(usuario)
        db.session.commit()
        flash(f'La cuenta de {usuario.nombre} fue eliminada.', 'success')
    return redirect(url_for('usuarios'))


@app.route('/solicitudes-recuperacion/<int:id>/atender', methods=['POST'])
@roles_requeridos('Administrador')
def atender_solicitud_recuperacion(id):
    solicitud = SolicitudRecuperacion.query.get_or_404(id)
    usuario = Usuario.query.filter_by(correo=solicitud.correo).first()
    nueva_password = request.form['nueva_password']
    if usuario and len(nueva_password) >= 8:
        usuario.password_hash = generate_password_hash(nueva_password)
        solicitud.estado = 'Atendida'
        db.session.commit()
        flash(f'Contraseña restablecida para {usuario.correo}.', 'success')
    else:
        flash('La contraseña debe tener al menos 8 caracteres.', 'danger')
    return redirect(url_for('usuarios'))

@app.route('/administracion')
@roles_requeridos('Administrador')
def inicio():
    pagina = request.args.get('pagina', 1, type=int)
    clientes = Cliente.query.order_by(Cliente.nombre_razon_social).paginate(page=pagina, per_page=25, error_out=False)
    return render_template('clientes.html', clientes=clientes)


@app.route('/nuevo_cliente', methods=['GET'])
@roles_requeridos('Administrador')
def mostrar_formulario_cliente():
    return render_template('nuevo_cliente.html', cliente=None,
                           usuarios_cliente=Usuario.query.filter_by(rol='Cliente').order_by(Usuario.nombre).all())


@app.route('/nuevo_cliente', methods=['POST'])
@roles_requeridos('Administrador')
def guardar_cliente():
    nuevo = Cliente(
        tipo_persona=request.form['tipo_persona'],
        numero_identificacion=request.form['numero_identificacion'],
        nombre_razon_social=request.form['nombre_razon_social'],
        telefono=request.form['telefono'],
        correo=request.form['correo'],
        id_usuario=request.form.get('id_usuario', type=int) or None
    )
    db.session.add(nuevo)
    db.session.commit()
    return redirect('/administracion')


@app.route('/editar_cliente/<int:id>', methods=['GET'])
@roles_requeridos('Administrador')
def mostrar_formulario_editar_cliente(id):
    cliente = Cliente.query.get_or_404(id)
    return render_template('nuevo_cliente.html', cliente=cliente, usuarios_cliente=Usuario.query.filter_by(rol='Cliente').order_by(Usuario.nombre).all())


@app.route('/editar_cliente/<int:id>', methods=['POST'])
@roles_requeridos('Administrador')
def actualizar_cliente(id):
    cliente = Cliente.query.get_or_404(id)
    cliente.tipo_persona = request.form['tipo_persona']
    cliente.numero_identificacion = request.form['numero_identificacion']
    cliente.nombre_razon_social = request.form['nombre_razon_social']
    cliente.telefono = request.form['telefono']
    cliente.correo = request.form['correo']
    cliente.id_usuario = request.form.get('id_usuario', type=int) or None
    db.session.commit()
    return redirect('/administracion')


@app.route('/eliminar_cliente/<int:id>', methods=['POST'])
@roles_requeridos('Administrador')
def eliminar_cliente(id):
    cliente = Cliente.query.get_or_404(id)
    app.logger.info('Eliminación de cliente: cliente=%s, administrador=%s', cliente.id_cliente, session['usuario_id'])
    db.session.delete(cliente)
    db.session.commit()
    return redirect('/administracion')

# ---------- RUTAS: PRODUCTOS ----------

@app.route('/productos')
@roles_requeridos('Administrador')
def ver_productos():
    pagina = request.args.get('pagina', 1, type=int)
    productos = Producto.query.order_by(Producto.tipo_estilo).paginate(page=pagina, per_page=25, error_out=False)
    return render_template('productos.html', productos=productos)


@app.route('/nuevo_producto', methods=['GET'])
@roles_requeridos('Administrador')
def mostrar_formulario_producto():
    return render_template('nuevo_producto.html', producto=None)


@app.route('/nuevo_producto', methods=['POST'])
@roles_requeridos('Administrador')
def guardar_producto():
    desarrollo_valor = request.form['desarrollo']
    if desarrollo_valor == '':
        desarrollo_valor = None

    nuevo = Producto(
        tipo_estilo=request.form['tipo_estilo'],
        calibre=request.form['calibre'],
        desarrollo=desarrollo_valor,
        precio_por_metro=request.form['precio_por_metro']
    )
    db.session.add(nuevo)
    db.session.commit()
    return redirect('/productos')


@app.route('/editar_producto/<int:id>', methods=['GET'])
@roles_requeridos('Administrador')
def mostrar_formulario_editar_producto(id):
    producto = Producto.query.get_or_404(id)
    return render_template('nuevo_producto.html', producto=producto)


@app.route('/editar_producto/<int:id>', methods=['POST'])
@roles_requeridos('Administrador')
def actualizar_producto(id):
    producto = Producto.query.get_or_404(id)
    desarrollo_valor = request.form['desarrollo']
    if desarrollo_valor == '':
        desarrollo_valor = None

    producto.tipo_estilo = request.form['tipo_estilo']
    producto.calibre = request.form['calibre']
    producto.desarrollo = desarrollo_valor
    producto.precio_por_metro = request.form['precio_por_metro']
    db.session.commit()
    return redirect('/productos')


@app.route('/eliminar_producto/<int:id>', methods=['POST'])
@roles_requeridos('Administrador')
def eliminar_producto(id):
    producto = Producto.query.get_or_404(id)
    app.logger.info('Eliminación de producto: producto=%s, administrador=%s', producto.id_producto, session['usuario_id'])
    db.session.delete(producto)
    db.session.commit()
    return redirect('/productos')

# ---------- RUTAS: PEDIDOS ----------

@app.route('/pedidos')
@roles_requeridos('Administrador', 'Operario', 'Vendedor')
def ver_pedidos():
    rol = session['usuario_rol']
    consulta = Pedido.query.order_by(Pedido.fecha_pedido.desc(), Pedido.id_pedido.desc())
    if rol == 'Vendedor':
        consulta = consulta.filter_by(id_vendedor=session['usuario_id'])
    periodo = request.args.get('periodo', 'todos')
    hoy = date.today()
    etiquetas_periodo = {
        'todos': 'Todos los pedidos',
        'hoy': 'Pedidos de hoy',
        'semana': 'Pedidos de esta semana',
        'mes': 'Pedidos de este mes',
        'registrados_hoy': 'Pedidos registrados hoy',
    }
    if periodo == 'hoy':
        consulta = consulta.filter(Pedido.fecha_pedido == hoy)
    elif periodo == 'semana':
        consulta = consulta.filter(Pedido.fecha_pedido >= hoy - timedelta(days=hoy.weekday()))
    elif periodo == 'mes':
        consulta = consulta.filter(Pedido.fecha_pedido >= hoy.replace(day=1))
    elif periodo == 'registrados_hoy':
        consulta = consulta.filter(func.date(Pedido.fecha_registro) == hoy)
    else:
        periodo = 'todos'
    pagina = request.args.get('pagina', 1, type=int)
    pedidos = consulta.paginate(page=pagina, per_page=25, error_out=False)

    resumen = None
    if rol == 'Vendedor':
        hoy = date.today()
        def ventas_desde(inicio):
            return db.session.query(func.count(Pedido.id_pedido), func.coalesce(func.sum(Pedido.total_pedido), 0)).filter(Pedido.id_vendedor == session['usuario_id'], Pedido.fecha_pedido >= inicio).one()
        semana = ventas_desde(hoy - timedelta(days=6))
        quincena = ventas_desde(hoy - timedelta(days=14))
        mes = ventas_desde(hoy.replace(day=1))
        resumen = {
            'Esta semana': {'pedidos': semana[0], 'total': semana[1]},
            'Últimos 15 días': {'pedidos': quincena[0], 'total': quincena[1]},
            'Este mes': {'pedidos': mes[0], 'total': mes[1]},
        }
    clientes_filtro = Cliente.query.join(Pedido).filter(Pedido.id_vendedor == session['usuario_id']).distinct().order_by(Cliente.nombre_razon_social).all() if rol == 'Vendedor' else Cliente.query.order_by(Cliente.nombre_razon_social).all()
    return render_template('pedidos.html', pedidos=pedidos, resumen=resumen, periodo=periodo,
                           etiqueta_periodo=etiquetas_periodo[periodo], clientes_filtro=clientes_filtro)


@app.route('/ver_pedido/<int:id>')
@roles_requeridos('Administrador', 'Operario', 'Vendedor')
def ver_pedido(id):
    pedido = Pedido.query.get_or_404(id)
    if session['usuario_rol'] == 'Vendedor' and pedido.id_vendedor != session['usuario_id']:
        flash('Solo puedes consultar los pedidos que registraste.', 'danger')
        return redirect(url_for('ver_pedidos'))
    return render_template('remision.html', pedido=pedido)


@app.route('/pedidos/<int:id>/estado', methods=['POST'])
@roles_requeridos('Operario')
def actualizar_estado_pedido(id):
    pedido = Pedido.query.get_or_404(id)
    nuevo_estado = request.form.get('estado')
    estados_validos = ('Pendiente', 'Listo para entrega', 'Entregado')

    if nuevo_estado not in estados_validos:
        flash('El estado seleccionado no es válido.', 'danger')
    elif nuevo_estado == 'Listo para entrega' and pedido.estado == 'Pendiente':
        faltantes = []
        for detalle in pedido.detalles:
            item = Inventario.query.filter_by(id_producto=detalle.id_producto, color=detalle.color).first() or Inventario.query.filter_by(id_producto=detalle.id_producto, color='General').first()
            if not item or item.cantidad_actual < detalle.cantidad:
                faltantes.append(detalle.producto.tipo_estilo)
        if faltantes:
            flash('No hay existencias suficientes para producción: ' + ', '.join(faltantes) + '.', 'danger')
            return redirect(url_for('ver_pedidos'))
        for detalle in pedido.detalles:
            item = Inventario.query.filter_by(id_producto=detalle.id_producto, color=detalle.color).first() or Inventario.query.filter_by(id_producto=detalle.id_producto, color='General').first()
            item.cantidad_actual -= detalle.cantidad
            db.session.add(MovimientoInventario(id_inventario=item.id_inventario, id_usuario=session['usuario_id'], tipo='Salida', cantidad=detalle.cantidad, motivo=f'Producción pedido #{pedido.id_pedido}'))
        pedido.estado = nuevo_estado
        registrar_historial(pedido, 'Pedido enviado a despacho. Inventario descontado.')
        db.session.commit()
        flash(f'El pedido #{pedido.id_pedido} ahora está: {nuevo_estado}.', 'success')
    else:
        pedido.estado = nuevo_estado
        registrar_historial(pedido, f'Estado actualizado a {nuevo_estado}')
        db.session.commit()
        flash(f'El pedido #{pedido.id_pedido} ahora está: {nuevo_estado}.', 'success')
    return redirect(url_for('ver_pedidos'))

@app.route('/nuevo_pedido', methods=['GET'])
@roles_requeridos('Administrador', 'Vendedor')
def mostrar_formulario_pedido():
    clientes = Cliente.query.all()
    productos = Producto.query.all()
    vendedores = Usuario.query.filter_by(rol='Vendedor').order_by(Usuario.nombre).all()
    return render_template('nuevo_pedido.html', clientes=clientes, productos=productos, vendedores=vendedores, pedido=None)


@app.route('/nuevo_pedido', methods=['POST'])
@roles_requeridos('Administrador', 'Vendedor')
def guardar_pedido():
    id_vendedor = session['usuario_id'] if session['usuario_rol'] == 'Vendedor' else request.form.get('id_vendedor', type=int)
    if not id_vendedor or not Usuario.query.filter_by(id_usuario=id_vendedor, rol='Vendedor').first():
        flash('Debes asignar un vendedor válido al pedido.', 'danger')
        return redirect(url_for('mostrar_formulario_pedido'))
    productos_ids = request.form.getlist('id_producto[]')
    medidas = request.form.getlist('medida_metros[]')
    cantidades = request.form.getlist('cantidad[]')
    colores = request.form.getlist('color[]')
    if not productos_ids or not (len(productos_ids) == len(medidas) == len(cantidades) == len(colores)):
        flash('Agrega al menos un producto y completa todos sus datos.', 'danger')
        return redirect(url_for('mostrar_formulario_pedido'))

    try:
        nuevo_pedido = Pedido(id_cliente=request.form['id_cliente'], id_vendedor=id_vendedor,
                              fecha_pedido=request.form['fecha_pedido'], estado='Pendiente')
        db.session.add(nuevo_pedido)
        db.session.flush()
        total = 0
        for id_producto, medida, cantidad, color in zip(productos_ids, medidas, cantidades, colores):
            producto = db.session.get(Producto, int(id_producto))
            medida = float(medida)
            cantidad = int(cantidad)
            if not producto or medida <= 0 or cantidad <= 0 or not color.strip():
                raise ValueError
            subtotal = medida * float(producto.precio_por_metro) * cantidad
            db.session.add(DetallePedido(id_pedido=nuevo_pedido.id_pedido, id_producto=producto.id_producto,
                                         color=color.strip(), medida_metros=medida, cantidad=cantidad, subtotal=subtotal))
            total += subtotal
        nuevo_pedido.total_pedido = total
        registrar_historial(nuevo_pedido, 'Pedido registrado')
        db.session.commit()
    except (ValueError, TypeError):
        db.session.rollback()
        flash('Revisa los productos, colores, medidas y cantidades del pedido.', 'danger')
        return redirect(url_for('mostrar_formulario_pedido'))

    return redirect('/pedidos')


@app.route('/editar_pedido/<int:id>', methods=['GET', 'POST'])
@roles_requeridos('Administrador', 'Vendedor')
def editar_pedido(id):
    pedido = Pedido.query.get_or_404(id)
    if session['usuario_rol'] == 'Vendedor' and pedido.id_vendedor != session['usuario_id']:
        flash('Solo puedes modificar los pedidos que registraste.', 'danger')
        return redirect(url_for('ver_pedidos'))

    if request.method == 'POST':
        pedido.id_cliente = request.form['id_cliente']
        pedido.fecha_pedido = request.form['fecha_pedido']
        if session['usuario_rol'] == 'Administrador':
            id_vendedor = request.form.get('id_vendedor', type=int)
            if not id_vendedor or not Usuario.query.filter_by(id_usuario=id_vendedor, rol='Vendedor').first():
                flash('Selecciona un vendedor válido.', 'danger')
                return redirect(url_for('editar_pedido', id=id))
            pedido.id_vendedor = id_vendedor
        productos_ids = request.form.getlist('id_producto[]')
        medidas = request.form.getlist('medida_metros[]')
        cantidades = request.form.getlist('cantidad[]')
        colores = request.form.getlist('color[]')
        if not productos_ids or not (len(productos_ids) == len(medidas) == len(cantidades) == len(colores)):
            flash('Agrega al menos un producto y completa todos sus datos.', 'danger')
            return redirect(url_for('editar_pedido', id=id))
        try:
            for detalle in pedido.detalles:
                db.session.delete(detalle)
            db.session.flush()
            total = 0
            for id_producto, medida, cantidad, color in zip(productos_ids, medidas, cantidades, colores):
                producto = db.session.get(Producto, int(id_producto))
                medida = float(medida)
                cantidad = int(cantidad)
                if not producto or medida <= 0 or cantidad <= 0 or not color.strip():
                    raise ValueError
                subtotal = medida * float(producto.precio_por_metro) * cantidad
                db.session.add(DetallePedido(id_pedido=pedido.id_pedido, id_producto=producto.id_producto,
                                             color=color.strip(), medida_metros=medida, cantidad=cantidad, subtotal=subtotal))
                total += subtotal
            pedido.total_pedido = total
            registrar_historial(pedido, 'Pedido actualizado')
            db.session.commit()
        except (ValueError, TypeError):
            db.session.rollback()
            flash('Revisa los productos, colores, medidas y cantidades del pedido.', 'danger')
            return redirect(url_for('editar_pedido', id=id))
        flash(f'Pedido #{pedido.id_pedido} actualizado correctamente.', 'success')
        return redirect(url_for('ver_pedidos'))

    clientes = Cliente.query.all()
    productos = Producto.query.all()
    vendedores = Usuario.query.filter_by(rol='Vendedor').order_by(Usuario.nombre).all()
    return render_template('nuevo_pedido.html', clientes=clientes, productos=productos, vendedores=vendedores, pedido=pedido)

if __name__ == '__main__':
    app.run(debug=os.environ.get('FLASK_DEBUG') == '1')
